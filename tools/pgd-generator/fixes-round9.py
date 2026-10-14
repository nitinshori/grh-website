#!/usr/bin/env python3
"""
fixes-round9.py  --  training before first use, on every live PGD; and the
Period Delay v011 Appendix 1 defect

30 September 2026, Nitin (Medical Director): every PGD must say that the
practitioner should complete the Get Real Health training for it, or
equivalent training, before first working under it, with the date recorded,
but training must not be a condition of access, because many pharmacists
bring experience of the same service from elsewhere. The clause is added by
grh_reissue.service_training_requirement (so every future reissue keeps it)
directly after the generic PGD training bullet.

Period Delay only: v011 (24 Sep 2026) was meant to change Appendix 1
question 5 (long journey) alone, but its edit anchored on "YES: exclude and
refer.", which every row shared, so questions 1, 2, 3, 6, 7 and 8 (previous
clot, stroke or heart attack; family history; smoking; surgery; immobility;
cancer) were also rewritten to "not an exclusion on its own". The body of
v011 and its change note still exclude them. v012 restores those six rows
and corrects the checklist heading, found by the adversarial training
review of 30 September 2026.

  python3 tools/pgd-generator/fixes-round9.py            all documents
  python3 tools/pgd-generator/fixes-round9.py ed uti     named documents
  python3 tools/pgd-generator/fixes-round9.py --dry      build to /tmp only
"""

import json
import os
import re
import shutil
import sys

import docx

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from grh_reissue import reissue, walk, set_text  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(HERE))
PARENT = os.path.dirname(ROOT)
APPROVED = os.path.join(PARENT, "PGD Rewrite 2026", "02 Approved")
PUBLIC = os.path.join(ROOT, "public", "pgd-documents")
MANIFEST_TS = os.path.join(ROOT, "src", "lib", "pgd-document-manifest.ts")
R9 = os.path.join(HERE, "round9")
DATE = "30 September 2026"

SHARED = (
    "Training before first use, applied to every Get Real Health PGD on 30 September 2026: the staff section "
    "now states that each practitioner should complete the Get Real Health training for this PGD, or "
    "equivalent training elsewhere, before first working under it, must read the PGD in full before first use, "
    "and that the date of completion is recorded. Training is recorded; it is not a condition of access, and "
    "the employing organisation remains responsible for being satisfied of competence. No clinical content "
    "changes."
)

FALLBACK_CHANGE = (
    "This PGD's staff section did not carry the requirement, common to every other Get Real Health PGD, that "
    "users have completed training in working under Patient Group Directions and been assessed as competent to "
    "work under a PGD before first working under this one. It is added, with the training-before-first-use "
    "statement."
)

# docx basenames that differ from the published PDF name
DOCX_PREFIX = {
    "covid-2026-27": "covid", "flu-2026-27": "flu", "chest-service": "chest", "dental-bridging": "dental",
    "hair-loss": "alopecia", "skin-infection": "skin", "sti-testing": "chlamydia",
}
TWINS = {"folic-acid": "b12-injection", "alopecia-minoxidil": "hair-loss"}

PD_BAD = ("YES: not an exclusion on its own. Count it as one venous thromboembolism risk factor, give the journey "
          "precautions (move around at least hourly, keep well hydrated, avoid alcohol and sedatives, consider "
          "graduated compression stockings) and record that they were given. Where another counted risk factor "
          "is also present, weigh the risks and benefits and record the decision.")
PD_HEAD_OLD = "Ask all eight. Any single YES excludes. Record the answers, not just the outcome."
PD_HEAD_NEW = ("Ask all eight. A YES to question 1, 2, 3, 6, 7 or 8, or a BMI of 30 or above at question 4, "
               "excludes. A YES to question 5 (a long journey) does not exclude on its own: it is counted as one "
               "risk factor and the journey precautions are given. Record the answers, not just the outcome.")
PD_CHANGE = (
    "Appendix 1 corrected. Version 011 was meant to change question 5 (a journey of 4 hours or more) alone, "
    "but questions 1, 2, 3, 6, 7 and 8 (previous clot, stroke or heart attack; family history of clot or "
    "clotting disorder; smoking; surgery; immobility; cancer) were also changed in error to read \"not an "
    "exclusion on its own\". They contradicted the exclusion criteria and the version 011 change record, which "
    "both keep them as exclusions. Each is restored to \"YES: exclude and refer\". The checklist heading now says "
    "which answers exclude and that question 5 is counted, not excluded. Found by the Get Real Health "
    "adversarial training review of 30 September 2026."
)


def period_delay_hook(doc):
    hits = [p for p in walk(doc, doc.element.body) if (p.text or "").strip() == PD_BAD]
    if len(hits) != 7:
        raise SystemExit(f"period-delay: expected 7 defective Appendix 1 rows, found {len(hits)}")
    # Order in the document: questions 1, 2, 3, 5, 6, 7, 8. Keep question 5 (index 3).
    for i, p in enumerate(hits):
        if i != 3:
            set_text(p, "YES: exclude and refer.")
    heads = [p for p in walk(doc, doc.element.body) if (p.text or "").strip() == PD_HEAD_OLD]
    if len(heads) != 1:
        raise SystemExit(f"period-delay: checklist heading found {len(heads)} times")
    set_text(heads[0], PD_HEAD_NEW)
    print("  period-delay: Appendix 1 questions 1, 2, 3, 6, 7, 8 restored as exclusions; heading corrected")


def issue_date(path):
    d = docx.Document(path)
    for p in walk(d, d.element.body):
        m = re.match(r"^Patient Group Direction, version (\d{3}), issued (\d{1,2} \w+ \d{4})", (p.text or "").strip())
        if m:
            return m.group(1), m.group(2)
    raise SystemExit(f"no strapline in {path}")


def masters():
    s = open(MANIFEST_TS).read()
    block = s.split("PGD_MASTER_FILES")[1].split("};")[0]
    return dict(re.findall(r'"([^"]+)":\s*"([^"]+)"', block))


def source_docx(pdf):
    base = pdf[:-4]
    m = re.match(r"^(.*)-(v\d{3})$", base)
    stem, ver = m.group(1), m.group(2)
    for cand in (f"{base}-SIGNED.docx", f"{DOCX_PREFIX.get(stem, stem)}-{ver}-SIGNED.docx"):
        p = os.path.join(APPROVED, cand)
        if os.path.exists(p):
            return p, stem, ver
    raise SystemExit(f"no source docx for {pdf}")


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    dry = "--dry" in sys.argv
    outdir = "/tmp/r9out" if dry else APPROVED
    os.makedirs(outdir, exist_ok=True)
    os.makedirs(R9, exist_ok=True)
    live = masters()
    manifest_path = os.path.join(R9, "manifest_out.json")
    manifest = json.load(open(manifest_path)) if os.path.exists(manifest_path) else {}
    for slug, pdf in sorted(live.items()):
        if slug in TWINS or (args and slug not in args):
            continue
        if slug in manifest and not dry:
            print(f"{slug:24} already reissued as {manifest[slug]}, skipped")
            continue
        src, stem, ver = source_docx(pdf)
        v_src, prev_date = issue_date(src)
        if v_src != ver[1:]:
            raise SystemExit(f"{slug}: docx strapline says {v_src}, file name says {ver}")
        num = int(ver[1:])
        version = f"v{num + 1:03d}"
        out_name = re.sub(r"-v\d{3}-SIGNED\.docx$", f"-{version}-SIGNED.docx", os.path.basename(src))
        out = os.path.join(outdir, out_name)
        pdf_name = f"{stem}-{version}.pdf"
        changes = [SHARED]
        hooks = ()
        src_text = "\n".join(p.text or "" for p in walk(docx.Document(src), docx.Document(src).element.body))
        if "assessed as competent to work under a PGD before first working under this one" not in src_text:
            changes = [SHARED, FALLBACK_CHANGE]
        if slug == "period-delay":
            changes = [PD_CHANGE, SHARED]
            hooks = (period_delay_hook,)
        built = reissue(slug, src, out, version, f"Version {num:03d}, {prev_date}", DATE, changes, [],
                        publish_to=None if dry else os.path.join(PUBLIC, pdf_name), remove=None, hooks=hooks)
        if dry:
            continue
        manifest[slug] = pdf_name
        for twin, of in TWINS.items():
            if of == slug:
                twin_pdf = pdf_name if twin == "folic-acid" else f"{twin}-{version}.pdf"
                if twin_pdf != pdf_name:
                    shutil.copyfile(built, os.path.join(PUBLIC, twin_pdf))
                manifest[twin] = twin_pdf
        json.dump(manifest, open(manifest_path, "w"), indent=1, sort_keys=True)
    print(f"\n{len(manifest)} manifest entries in {manifest_path}")


if __name__ == "__main__":
    main()
